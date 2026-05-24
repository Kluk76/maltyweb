"""
ingest_bd_racking_v2.py — upload normalized RackingData into bd_racking_v2.

Source: /var/www/maltytask/data/RawDB-normalized.xlsx, sheet RackingData (399 rows).

Column map (normalized sheet, 0-based):
  0  Timestamp            → submitted_at (datetime OR "MM/DD/YYYY HH:MM:SS" string)
  1  Email Address        → email
  2  Date du dernier CIP  → last_cip_date (string, kept verbatim)
  3  Type de CIP          → cip_type
  4  Type de Rack         → rack_type
  5  Client               → client
  6  Recettes Nébuleuse   → neb_beer
  7  Batch                → neb_batch
  8  Recettes Contract    → contract_beer
  9  Batch                → contract_batch
  10 Start Time           → start_time (combined with event_date)
  11 End Time             → end_time   (combined with event_date)
  12 BBT                  → bbt_old (legacy bare-int; fallback destination only)
  13 CO2 in BBT           → bbt_co2 (datetime/text → NULL + flag)
  14 O2 in BBT            → bbt_o2
  15 Total Racked Vol     → racked_vol_hl
  16 Blend                → blend_hl (non-numeric → NULL + flag)
  17 Average Turbidity    → avg_turbidity
  18 Average Speed        → avg_speed
  19 Pressure in BBT      → bbt_pressure
  20 Centri Rinsed ?      → centri_rinsed
  21 Final Comments       → comments
  22 CIP BBT ?            → cip_bbt_done
  23 CIP Type             → cip_bbt_type
  24 Date of BBT CIP      → cip_bbt_date
  25 Which BBT            → CANONICAL destination ("BBT 1".."CCT 5")
  26 neb_recipe_id        → neb_recipe_id_fk
  27 contract_recipe_id   → contract_recipe_id_fk

Usage:
  python3 ingest_bd_racking_v2.py            # dry-run
  python3 ingest_bd_racking_v2.py --apply    # live
  python3 ingest_bd_racking_v2.py --apply --verify
  python3 ingest_bd_racking_v2.py --limit 20
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, date, time
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

_SCRIPT_DIR = Path(__file__).parent
if str(_SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPT_DIR))

from lib_config import load as load_config
from lib_db import connect

XLSX_PATH    = Path("/var/www/maltytask/data/RawDB-normalized.xlsx")
SHEET_NAME   = "RackingData"
TARGET_TABLE = "bd_racking_v2"

NK_GROUP = ("submitted_at", "neb_beer", "neb_batch", "contract_beer", "contract_batch")
_META_KEYS = frozenset({"row_hash", "is_tombstoned", "audit_flags",
                        "imported_at", "updated_at", "id"})

# ── Type coercions ──────────────────────────────────────────────────────────────

def _s(v: Any) -> str | None:
    if v is None: return None
    s = str(v).strip()
    return s if s else None

def _i(v: Any) -> int | None:
    if v is None: return None
    if isinstance(v, bool): return int(v)
    if isinstance(v, (int, float)): return int(v)
    s = str(v).strip()
    if not s: return None
    try: return int(float(s.replace(",", ".")))
    except (ValueError, TypeError): return None

def _f(v: Any) -> float | None:
    """Float coercion. European decimal comma tolerated (3,91 → 3.91).
    datetime/time/unparseable → None (caller flags)."""
    if v is None: return None
    if isinstance(v, bool): return float(v)
    if isinstance(v, (int, float)): return float(v)
    if isinstance(v, (datetime, date, time)): return None  # type anomaly
    s = str(v).strip()
    if not s: return None
    try: return float(s.replace(",", "."))
    except (ValueError, TypeError): return None

def _dt(v: Any) -> str | None:
    """MySQL DATETIME(6). Handles datetime objects AND 'MM/DD/YYYY HH:MM:SS' strings."""
    if v is None: return None
    if isinstance(v, datetime): return v.strftime("%Y-%m-%d %H:%M:%S.%f")
    if isinstance(v, date): return datetime(v.year, v.month, v.day).strftime("%Y-%m-%d %H:%M:%S.000000")
    s = str(v).strip()
    if not s: return None
    for fmt in ("%m/%d/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y", "%Y-%m-%d"):
        try: return datetime.strptime(s, fmt).strftime("%Y-%m-%d %H:%M:%S.%f")
        except ValueError: pass
    return None

def _date_of(dt_str: str | None) -> str | None:
    if not dt_str: return None
    try: return datetime.strptime(dt_str, "%Y-%m-%d %H:%M:%S.%f").strftime("%Y-%m-%d")
    except ValueError: return None

def _combine(dt_str: str | None, t: Any) -> str | None:
    """Combine the date part of submitted_at with a time-of-day → DATETIME string."""
    d = _date_of(dt_str)
    if not d or t is None: return None
    if isinstance(t, time):
        return f"{d} {t.hour:02d}:{t.minute:02d}:{t.second:02d}"
    if isinstance(t, datetime):
        return f"{d} {t.hour:02d}:{t.minute:02d}:{t.second:02d}"
    return None

def _row_hash(d: dict) -> str:
    return hashlib.sha256(json.dumps(d, sort_keys=True, default=str).encode("utf-8")).hexdigest()

def _is_empty(row: tuple) -> bool:
    return not any(v is not None and str(v).strip() != "" for v in row)


# ── Destination tank resolution ─────────────────────────────────────────────────

_TANK_RE = re.compile(r"^\s*([A-Za-z]+)\s*0*(\d+)\s*$")

def resolve_destination(c25: Any, c12: Any, valid_bbt: set, valid_cct: set) -> dict:
    """
    Resolve the racking destination tank.
    Canonical = "Which BBT" (col 25) as "BBT N"/"CCT N"/"YT N".
    Fallback   = legacy bare-int "BBT" (col 12) → treated as BBT N.
    Returns {dest_type, bbt_number, cct_number, target_tank_raw, flags[]}.
    """
    flags: list[str] = []
    raw25 = _s(c25)
    bbt_old = _i(c12)

    dest_type = bbt_number = cct_number = None
    target_raw = None

    if raw25:
        target_raw = raw25
        m = _TANK_RE.match(raw25)
        if m:
            prefix = m.group(1).upper()
            num = int(m.group(2))
            if prefix == "BBT":
                dest_type, bbt_number = "BBT", num
            elif prefix == "CCT":
                dest_type, cct_number = "CCT", num
            elif prefix in ("YT", "YEAST"):
                dest_type = "YT"  # no FK column; recorded via dest_type only
            else:
                flags.append(f"dest_unparsed:{raw25}")
        else:
            flags.append(f"dest_unparsed:{raw25}")
    elif bbt_old is not None:
        dest_type, bbt_number = "BBT", bbt_old
        target_raw = f"BBT {bbt_old}"
    # else: no destination given (3 rows) — all NULL, no flag (legitimately absent)

    # Validate against ref tables; out-of-range → NULL FK + flag (keeps row insertable)
    if bbt_number is not None and bbt_number not in valid_bbt:
        flags.append(f"bbt_not_in_ref:{bbt_number}")
        bbt_number = None
    if cct_number is not None and cct_number not in valid_cct:
        flags.append(f"cct_not_in_ref:{cct_number}")
        cct_number = None

    return {
        "racking_destination_type": dest_type,
        "bbt_number": bbt_number,
        "cct_number": cct_number,
        "target_tank_raw": target_raw,
        "bbt_old": bbt_old,
        "_flags": flags,
    }


# ── Lookup helpers ──────────────────────────────────────────────────────────────

def load_recipe_set(conn) -> set[int]:
    with conn.cursor() as cur:
        cur.execute("SELECT id FROM ref_recipes")
        return {r["id"] for r in cur.fetchall()}

def load_valid_bbt(conn) -> set[int]:
    with conn.cursor() as cur:
        cur.execute("SELECT number FROM ref_bbt")
        return {r["number"] for r in cur.fetchall()}

def load_valid_cct(conn) -> set[int]:
    with conn.cursor() as cur:
        cur.execute("SELECT number FROM ref_cct")
        return {r["number"] for r in cur.fetchall()}


# ── Parse ───────────────────────────────────────────────────────────────────────

def parse_rows(ws, recipe_set: set, valid_bbt: set, valid_cct: set, limit: int | None) -> list[dict]:
    out: list[dict] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if _is_empty(row): continue
        if limit is not None and len(out) >= limit: break
        r = list(row) + [None] * max(0, 28 - len(row))

        submitted_at = _dt(r[0])
        if not submitted_at:
            continue  # submitted_at is part of NK — skip if unparseable
        event_date = _date_of(submitted_at)

        neb_beer  = _s(r[6]) or ""
        neb_batch = str(_i(r[7])) if r[7] is not None else ""
        con_beer  = _s(r[8]) or ""
        con_batch = str(_i(r[9])) if r[9] is not None else ""

        flags: list[str] = []
        if not neb_beer and not con_beer:
            flags.append("no_beer_identity")

        # recipe FKs (already resolved by normalizer)
        neb_rid = _i(r[26])
        if neb_rid is not None and neb_rid not in recipe_set:
            flags.append(f"neb_recipe_not_found:{neb_rid}"); neb_rid = None
        con_rid = _i(r[27])
        if con_rid is not None and con_rid not in recipe_set:
            flags.append(f"contract_recipe_not_found:{con_rid}"); con_rid = None

        dest = resolve_destination(r[25], r[12], valid_bbt, valid_cct)
        flags.extend(dest["_flags"])

        bbt_co2 = _f(r[13])
        if bbt_co2 is None and r[13] is not None and str(r[13]).strip():
            flags.append("co2_unparseable")
        bbt_o2 = _f(r[14])
        if bbt_o2 is None and r[14] is not None and str(r[14]).strip():
            flags.append("o2_unparseable")
        blend_hl = _f(r[16])
        if blend_hl is None and r[16] is not None and str(r[16]).strip():
            flags.append("blend_nonnumeric")

        rowd = {
            "submitted_at": submitted_at,
            "email": _s(r[1]),
            "event_date": event_date,
            "seq": 0,  # set in _finalize
            "neb_beer": neb_beer,
            "neb_batch": neb_batch,
            "neb_recipe_id_fk": neb_rid,
            "contract_beer": con_beer,
            "contract_batch": con_batch,
            "contract_recipe_id_fk": con_rid,
            "last_cip_date": _s(r[2]),
            "cip_type": _s(r[3]),
            "rack_type": _s(r[4]),
            "client": _s(r[5]),
            "start_time": _combine(submitted_at, r[10]),
            "end_time": _combine(submitted_at, r[11]),
            "racking_destination_type": dest["racking_destination_type"],
            "bbt_number": dest["bbt_number"],
            "cct_number": dest["cct_number"],
            "target_tank_raw": dest["target_tank_raw"],
            "bbt_old": dest["bbt_old"],
            "bbt_co2": bbt_co2,
            "bbt_o2": bbt_o2,
            "racked_vol_hl": _f(r[15]),
            "blend_hl": blend_hl,
            "avg_turbidity": _f(r[17]),
            "avg_speed": _f(r[18]),
            "bbt_pressure": _f(r[19]),
            "centri_rinsed": _s(r[20]),
            "comments": _s(r[21]),
            "cip_bbt_done": _s(r[22]),
            "cip_bbt_type": _s(r[23]),
            "cip_bbt_date": _s(r[24]),
            "is_tombstoned": 0,
            "audit_flags": ",".join(flags) if flags else None,
        }
        out.append(rowd)
    return out


def _finalize(rows: list[dict]) -> list[dict]:
    """Assign content-sorted seq within each NK group, then compute row_hash.
    Handles the rare same-second multi-racking collision deterministically."""
    groups: dict[tuple, list[dict]] = defaultdict(list)
    for row in rows:
        groups[tuple(row[k] for k in NK_GROUP)].append(row)
    for key, group in groups.items():
        if len(group) > 1:
            group.sort(key=lambda r: (
                str(r["rack_type"] or ""),
                r["bbt_number"] if r["bbt_number"] is not None else -1,
                r["cct_number"] if r["cct_number"] is not None else -1,
                str(r["start_time"] or ""),
                str(r["comments"] or ""),
            ))
            for idx, row in enumerate(group):
                row["seq"] = idx
    for row in rows:
        canonical = {k: v for k, v in row.items() if k not in _META_KEYS}
        row["row_hash"] = _row_hash(canonical)
    return rows


# ── Upsert ──────────────────────────────────────────────────────────────────────

def upsert_rows(conn, rows: list[dict], dry_run: bool) -> None:
    if not rows:
        print("  0 rows — nothing to upsert"); return
    cols = list(rows[0].keys())
    placeholders = ", ".join(["%s"] * len(cols))
    col_names = ", ".join(f"`{c}`" for c in cols)
    nk = set(NK_GROUP) | {"seq"}
    update_cols = [c for c in cols if c not in nk and c not in ("row_hash", "imported_at", "id")]
    update_clause = ", ".join(f"`{c}`=VALUES(`{c}`)" for c in update_cols)
    sql = (f"INSERT INTO `{TARGET_TABLE}` ({col_names}) VALUES ({placeholders})\n"
           f"ON DUPLICATE KEY UPDATE {update_clause}")
    if dry_run:
        print(f"  DRY-RUN: would upsert {len(rows)} rows")
        for r in rows[:3]:
            print(f"    {json.dumps({k: str(v)[:32] for k, v in r.items() if v is not None})}")
        return
    affected = 0
    with conn.cursor() as cur:
        for start in range(0, len(rows), 500):
            batch = rows[start:start + 500]
            cur.executemany(sql, [[r[c] for c in cols] for r in batch])
            affected += cur.rowcount
    conn.commit()
    print(f"  upserted {len(rows)} rows (MySQL affected: {affected})")


def run_verification(conn) -> None:
    checks = [
        ("total rows", "SELECT COUNT(*) AS n FROM bd_racking_v2"),
        ("destination types", "SELECT racking_destination_type AS t, COUNT(*) AS n FROM bd_racking_v2 GROUP BY racking_destination_type ORDER BY n DESC"),
        ("recipe NULL (no beer)", "SELECT COUNT(*) AS n FROM bd_racking_v2 WHERE neb_recipe_id_fk IS NULL AND contract_recipe_id_fk IS NULL"),
        ("seq>0 (multi-racking)", "SELECT COUNT(*) AS n FROM bd_racking_v2 WHERE seq > 0"),
        ("flagged rows", "SELECT audit_flags, COUNT(*) AS n FROM bd_racking_v2 WHERE audit_flags IS NOT NULL GROUP BY audit_flags"),
        ("start_time populated", "SELECT COUNT(*) AS n FROM bd_racking_v2 WHERE start_time IS NOT NULL"),
    ]
    print("\n=== Verification ===")
    with conn.cursor() as cur:
        for label, sql in checks:
            cur.execute(sql)
            print(f"\n  {label}:")
            for r in cur.fetchall():
                print(f"    {dict(r)}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--limit", type=int, default=None)
    ap.add_argument("--verify", action="store_true")
    args = ap.parse_args()
    dry_run = not args.apply

    if not XLSX_PATH.exists():
        print(f"ERROR: xlsx not found at {XLSX_PATH}", file=sys.stderr); sys.exit(1)

    cfg = load_config(); conn = connect(cfg)
    recipe_set = load_recipe_set(conn)
    valid_bbt = load_valid_bbt(conn)
    valid_cct = load_valid_cct(conn)
    print(f"{'DRY-RUN' if dry_run else 'LIVE'} — {SHEET_NAME}")
    print(f"Loaded: {len(recipe_set)} recipes, {len(valid_bbt)} BBTs, {len(valid_cct)} CCTs\n")

    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    rows = _finalize(parse_rows(wb[SHEET_NAME], recipe_set, valid_bbt, valid_cct, args.limit))
    wb.close()
    print(f"  parsed {len(rows)} rows")
    upsert_rows(conn, rows, dry_run)

    if not dry_run and args.verify:
        run_verification(conn)
    conn.close()
    print("Done.")


if __name__ == "__main__":
    main()
